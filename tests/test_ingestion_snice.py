"""Tests del parser SNICE (LIGIE/tarifa y NICO) — capítulos 84/85 (§ TAREA LIGIE/NICO).

`unit` — corre siempre, sobre filas construidas a mano y un XLSX diminuto que
imita el layout real verificado contra los archivos que publica SNICE.
"""

from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from ingestion.snice import nico, tariff

pytestmark = pytest.mark.unit

_META = {
    "source_url": "https://www.snice.gob.mx/~oracle/SNICE_DOCS/FRACCIONESARANCELARIAS-LIGIE_20260420-20260420.xlsx",
    "content_hash": "deadbeef",
    "retrieved_at": datetime(2026, 9, 7, tzinfo=UTC),
}

_CHAPTERS_84_85 = frozenset({"84", "85"})


def _fa_row(
    code: str, description: object, unit: str = "Kg", igi: object = 15, ige: object = "Ex."
):
    return (None, None, code, description, unit, igi, ige, None, None)


# ── parse_fracciones ─────────────────────────────────────────────────────────


def test_filtra_por_capitulo() -> None:
    rows = [_fa_row("8401.10.01", "Reactores nucleares."), _fa_row("0101.21.01", "Un caballo.")]

    result = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META)

    assert [f.code for f in result.accepted] == ["84011001"]


def test_parse_fracciones_pobla_specificity() -> None:
    rows = [_fa_row("8401.10.01", "Los demás.")]

    fraccion = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META).accepted[0]

    assert fraccion.specificity == 0


def test_deriva_capitulo_partida_subpartida_del_codigo() -> None:
    rows = [_fa_row("8401.10.01", "Reactores nucleares.")]

    fraccion = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META).accepted[0]

    assert (fraccion.chapter, fraccion.heading, fraccion.subheading) == ("84", "8401", "840110")


def test_rechaza_sin_descripcion_en_vez_de_inventarla() -> None:
    rows = [_fa_row("8401.10.01", None)]

    result = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META)

    assert result.accepted == []
    assert result.rejected[0].reason == "sin descripción en la fuente"


def test_rechaza_codigo_malformado() -> None:
    rows = [_fa_row("no-es-un-codigo", "algo")]

    result = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META)

    assert result.rejected[0].reason == "código no tiene forma NNNN.NN.NN"


def test_rechaza_unidad_que_no_cabe_en_varchar4() -> None:
    """Caso real de la fuente: unidad='Prohibida' (9 caracteres) cuando la operación está vedada."""
    rows = [
        _fa_row(
            "8543.40.01",
            "Cigarrillos electrónicos.",
            unit="Prohibida",
            igi="Prohibida",
            ige="Prohibida",
        )
    ]

    result = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META)

    assert result.accepted == []
    assert "VARCHAR(4)" in result.rejected[0].reason


def test_tasa_ex_se_convierte_a_cero() -> None:
    rows = [_fa_row("8401.10.01", "Reactores nucleares.", igi="Ex.", ige="Ex.")]

    fraccion = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META).accepted[0]

    assert fraccion.igi_rate == Decimal("0")
    assert fraccion.ige_rate == Decimal("0")


def test_tasa_numerica_se_convierte_a_fraccion_no_a_entero() -> None:
    rows = [_fa_row("8402.11.01", "Calderas.", igi=15, ige="Ex.")]

    fraccion = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META).accepted[0]

    assert fraccion.igi_rate == Decimal("15") / Decimal(100)


def test_tasa_no_reconocida_es_advertencia_no_rechazo() -> None:
    """La fracción tiene código y descripción válidos: se acepta con la tasa en None."""
    rows = [_fa_row("8401.10.01", "Reactores nucleares.", igi="ver nota 3", ige="Ex.")]

    result = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META)

    assert len(result.accepted) == 1
    assert result.accepted[0].igi_rate is None
    assert "no reconocida" in result.rate_warnings[0]


def test_source_document_distingue_el_instrumento_del_render() -> None:
    """Decisión de Persona 1: source_document cita la ley, source_url de dónde se leyó."""
    rows = [_fa_row("8401.10.01", "Reactores nucleares.")]

    fraccion = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META).accepted[0]

    assert fraccion.source_document == "LIGIE 2022 (DOF)"
    assert fraccion.source_url.startswith("https://www.snice.gob.mx")


def test_reconciliacion_cuadra_cuando_los_tres_numeros_coinciden() -> None:
    rows = [_fa_row("8401.10.01", "Reactores nucleares."), _fa_row("8401.20.01", None)]

    result = tariff.parse_fracciones(rows, _CHAPTERS_84_85, **_META)

    assert "cuadra=True" in result.reconciliation(declared_by_source=2)
    assert "cuadra=False" in result.reconciliation(declared_by_source=3)


def test_specificity_catch_all_es_la_mas_baja() -> None:
    assert tariff.specificity_of("Los demás.") == 0
    assert tariff.specificity_of("Las demás presentadas en forma de sistemas.") == 0


def test_specificity_exclusion_es_intermedia() -> None:
    texto = "Unidades de proceso, excepto las de las subpartidas 8471.41 u 8471.49."
    assert tariff.specificity_of(texto) == 1


def test_specificity_afirmativa_sin_numero() -> None:
    texto = (
        "Que incluyan en la misma envoltura, al menos, una unidad central de "
        "proceso y, aunque estén combinadas, una unidad de entrada y una de salida."
    )
    assert tariff.specificity_of(texto) == 2


def test_specificity_afirmativa_con_numero_es_la_mas_alta() -> None:
    texto = (
        "Máquinas automáticas para tratamiento o procesamiento de datos, "
        "portátiles, de peso inferior o igual a 10 kg, que estén constituidas, "
        "al menos, por una unidad central de proceso, un teclado y un visualizador."
    )
    assert tariff.specificity_of(texto) == 3


def test_specificity_desempata_el_caso_real_de_la_partida_8471() -> None:
    """Regresión: caso real reportado por Persona 1 — 5 subpartidas empatadas
    en 8471 porque specificity no existía. Con la heurística, el laptop
    (84713001) queda estrictamente por encima de las demás."""
    laptop = tariff.specificity_of(
        "Máquinas automáticas para tratamiento o procesamiento de datos, "
        "portátiles, de peso inferior o igual a 10 kg, que estén constituidas, "
        "al menos, por una unidad central de proceso, un teclado y un visualizador."
    )
    unidad_combinada = tariff.specificity_of(
        "Que incluyan en la misma envoltura, al menos, una unidad central de "
        "proceso y, aunque estén combinadas, una unidad de entrada y una de salida."
    )
    sistemas = tariff.specificity_of("Las demás presentadas en forma de sistemas.")
    unidades_proceso = tariff.specificity_of(
        "Unidades de proceso, excepto las de las subpartidas 8471.41 u 8471.49, "
        "aunque incluyan en la misma envoltura uno o dos de los tipos siguientes "
        "de unidades: unidad de memoria, unidad de entrada y unidad de salida."
    )
    otras_unidades = tariff.specificity_of(
        "Las demás unidades de máquinas automáticas para tratamiento o procesamiento de datos."
    )

    assert laptop > unidad_combinada
    assert laptop > sistemas
    assert laptop > unidades_proceso
    assert laptop > otras_unidades


def test_iter_fraccion_rows_lee_el_layout_real_de_dos_encabezados(tmp_path: Path) -> None:
    """Fila 9 en adelante: dos renglones de encabezado antes, verificado contra el archivo real."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FA"
    for _ in range(6):
        ws.append([None] * 9)
    ws.append([None, None, "Fracción Arancelaria", "Descripción", "Unidad de Medida", "Arancel %"])
    ws.append([None, None, None, None, None, "IMP.", "EXP."])
    ws.append([None, None, "8401.10.01", "Reactores nucleares.", "Kg", "Ex.", "Ex.", None, None])
    path = tmp_path / "fracciones.xlsx"
    wb.save(path)

    rows = list(tariff.iter_fraccion_rows(path))

    assert rows == [
        (None, None, "8401.10.01", "Reactores nucleares.", "Kg", "Ex.", "Ex.", None, None)
    ]


# ── parse_nicos ──────────────────────────────────────────────────────────────


def test_nico_filtra_por_capitulo_y_arma_full_code() -> None:
    rows = [("8401.10.01", "00", "Reactores nucleares."), ("0101.21.01", "00", "Un caballo.")]

    result = nico.parse_nicos(rows, _CHAPTERS_84_85, **_META)

    assert [n.full_code for n in result.accepted] == ["8401100100"]


def test_nico_rechaza_sin_codigo_nico() -> None:
    rows = [("8401.10.01", None, "Reactores nucleares.")]

    result = nico.parse_nicos(rows, _CHAPTERS_84_85, **_META)

    assert result.rejected[0].reason == "sin código NICO en la fuente"


def test_nico_rechaza_sin_descripcion() -> None:
    rows = [("8401.10.01", "00", None)]

    result = nico.parse_nicos(rows, _CHAPTERS_84_85, **_META)

    assert result.rejected[0].reason == "sin descripción en la fuente"


def test_nico_source_document_es_snice_no_dof() -> None:
    """Para NICO, SNICE es la fuente misma (decisión de Persona 1) — no LIGIE (DOF)."""
    rows = [("8401.10.01", "00", "Reactores nucleares.")]

    parsed = nico.parse_nicos(rows, _CHAPTERS_84_85, **_META).accepted[0]

    assert parsed.source_document == "NICO 2022 (SNICE)"


def test_iter_nico_rows_from_tariff_workbook_usa_columnas_d_e_f(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NICO"
    for _ in range(6):
        ws.append([None] * 6)
    ws.append([None, None, None, "FRACCIÓN ARANCELARIA", "NICO", "DESCRIPCIÓN"])
    ws.append([None, None, None, "8401.10.01", "00", "Reactores nucleares."])
    path = tmp_path / "nico_embebido.xlsx"
    wb.save(path)

    rows = list(nico.iter_nico_rows_from_tariff_workbook(path))

    assert rows == [("8401.10.01", "00", "Reactores nucleares.")]


def test_iter_nico_rows_from_standalone_usa_columnas_b_c_d(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NICO (ÚNICAMENTE)"
    for _ in range(4):
        ws.append([None] * 4)
    ws.append([None, "FRACCIÓN ARANCELARIA", "NICO", "DESCRIPCIÓN"])
    ws.append([None, "8401.10.01", "00", "Reactores nucleares."])
    path = tmp_path / "nico_standalone.xlsx"
    wb.save(path)

    rows = list(nico.iter_nico_rows_from_standalone(path))

    assert rows == [("8401.10.01", "00", "Reactores nucleares.")]
