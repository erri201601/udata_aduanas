"""PARSED -> NORMALIZED -> VALIDATED del NICO (Número de Identificación Comercial).

A diferencia de la tarifa, para NICO SNICE es la fuente misma, no un
complementario (decisión de Persona 1, 2026-09-05).

SNICE publica el mismo dato en dos archivos con layouts de columna
distintos: la hoja `NICO` embebida en el XLSX de fracciones, y el XLSX
`NICO (ÚNICAMENTE)` independiente. Los dos iteradores de abajo normalizan
cada uno a la misma tupla `(fracción, nico, descripción)` para que
`parse_nicos` no tenga que saber de cuál vinieron.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import TYPE_CHECKING

import openpyxl

if TYPE_CHECKING:
    from collections.abc import Iterator
    from datetime import datetime
    from pathlib import Path

SOURCE_DOCUMENT_NICO = "NICO 2022 (SNICE)"


@dataclass(frozen=True)
class ParsedNico:
    """Un NICO ya normalizado, listo para `Nico` (pendiente resolver `tariff_fraction_id`)."""

    fraction_code: str
    code: str
    full_code: str
    description: str
    source_document: str
    source_url: str
    content_hash: str
    retrieved_at: datetime


@dataclass(frozen=True)
class Rejection:
    raw_fraction: object
    reason: str


@dataclass(frozen=True)
class ParseResult:
    accepted: list[ParsedNico]
    rejected: list[Rejection]

    def reconciliation(self, *, declared_by_source: int) -> str:
        return (
            f"fuente={declared_by_source} "
            f"parseados={len(self.accepted)} "
            f"rechazados={len(self.rejected)} "
            f"(cuadra={declared_by_source == len(self.accepted) + len(self.rejected)})"
        )


def iter_nico_rows_from_tariff_workbook(
    path: str | Path,
) -> Iterator[tuple[object, object, object]]:
    """Hoja `NICO` embebida en el XLSX de fracciones (columnas D/E/F)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["NICO"]
    for row in ws.iter_rows(min_row=8, values_only=True):
        yield row[3], row[4], row[5]


def iter_nico_rows_from_standalone(path: str | Path) -> Iterator[tuple[object, object, object]]:
    """XLSX `NICO (ÚNICAMENTE)` independiente (columnas B/C/D)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["NICO (ÚNICAMENTE)"]
    for row in ws.iter_rows(min_row=6, values_only=True):
        yield row[1], row[2], row[3]


def _split_fraction(raw_code: str) -> str | None:
    digits = raw_code.replace(".", "").strip()
    if len(digits) != 8 or not digits.isdigit():
        return None
    return digits


def parse_nicos(
    rows: Iterator[tuple[object, object, object]],
    chapters: frozenset[str],
    *,
    source_url: str,
    content_hash: str,
    retrieved_at: datetime,
) -> ParseResult:
    accepted: list[ParsedNico] = []
    rejected: list[Rejection] = []

    for raw_fraction, raw_nico, description in rows:
        if raw_fraction is None:
            continue

        fraction_code = _split_fraction(str(raw_fraction))
        if fraction_code is None:
            rejected.append(Rejection(raw_fraction, "código de fracción no tiene forma NNNN.NN.NN"))
            continue
        if fraction_code[:2] not in chapters:
            continue

        nico_code = str(raw_nico).strip() if raw_nico is not None else ""
        if not nico_code:
            rejected.append(Rejection(raw_fraction, "sin código NICO en la fuente"))
            continue

        if description is None or not str(description).strip():
            rejected.append(Rejection(raw_fraction, "sin descripción en la fuente"))
            continue

        accepted.append(
            ParsedNico(
                fraction_code=fraction_code,
                code=nico_code,
                full_code=f"{fraction_code}{nico_code}",
                description=str(description).strip(),
                source_document=SOURCE_DOCUMENT_NICO,
                source_url=source_url,
                content_hash=content_hash,
                retrieved_at=retrieved_at,
            )
        )

    return ParseResult(accepted=accepted, rejected=rejected)
